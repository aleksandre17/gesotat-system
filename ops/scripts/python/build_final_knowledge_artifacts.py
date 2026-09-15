from pathlib import Path
from html import escape
from base64 import b64encode
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'documentation'/'complete-package'
OUT.mkdir(parents=True,exist_ok=True)

def inline(s):
    s=escape(s)
    s=re.sub(r'`([^`]+)`',r'<code>\1</code>',s)
    s=re.sub(r'\*\*([^*]+)\*\*',r'<strong>\1</strong>',s)
    return s

def markdown(text):
    lines=text.splitlines(); out=[]; i=0; code=False; buf=[]; ul=False; ol=False
    while i<len(lines):
        line=lines[i]
        if line.startswith('```'):
            if code: out.append('<pre><code>'+escape('\n'.join(buf))+'</code></pre>'); buf=[]; code=False
            else: code=True
            i+=1; continue
        if code: buf.append(line); i+=1; continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|?[\s:|-]+\|?$',lines[i+1]):
            rows=[]; heads=[x.strip() for x in line.strip('|').split('|')]; i+=2
            while i<len(lines) and lines[i].startswith('|'):
                rows.append([x.strip() for x in lines[i].strip('|').split('|')]); i+=1
            out.append('<table><thead><tr>'+''.join('<th>'+inline(x)+'</th>' for x in heads)+'</tr></thead><tbody>')
            for row in rows: out.append('<tr>'+''.join('<td>'+inline(x)+'</td>' for x in row)+'</tr>')
            out.append('</tbody></table>'); continue
        m=re.match(r'^(#{1,6})\s+(.+)$',line)
        if m:
            if ul: out.append('</ul>'); ul=False
            if ol: out.append('</ol>'); ol=False
            lvl=min(len(m.group(1)),6); out.append(f'<h{lvl}>{inline(m.group(2))}</h{lvl}>'); i+=1; continue
        if line.startswith('- '):
            if not ul: out.append('<ul>'); ul=True
            out.append('<li>'+inline(line[2:])+'</li>'); i+=1; continue
        if re.match(r'^\d+\. ',line):
            if not ol: out.append('<ol>'); ol=True
            out.append('<li>'+inline(re.sub(r'^\d+\. ','',line))+'</li>'); i+=1; continue
        if ul: out.append('</ul>'); ul=False
        if ol: out.append('</ol>'); ol=False
        if line.strip(): out.append('<p>'+inline(line.strip())+'</p>')
        i+=1
    if ul: out.append('</ul>')
    if ol: out.append('</ol>')
    return ''.join(out)

guide=(ROOT/'docs'/'system-complete-learning-guide.md').read_text(encoding='utf-8')
guide_html=markdown(guide)
images=[]
for n,c in [('01_planes.png','Control Data Archive planes'),('02_workflow.png','Workflow'),('03_generation.png','Contract generation'),('04_stat_model.png','Statistical model'),('05_runtime.png','Runtime execution')]:
    p=ROOT/'documentation'/'complete-package'/'source-assets'/n
    if p.exists() and p.stat().st_size:
        images.append('<figure><img src="data:image/png;base64,'+b64encode(p.read_bytes()).decode('ascii')+'" alt="'+c+'"><figcaption>'+c+'</figcaption></figure>')
status_names={'KIDS-R8-current-status-and-acceptance.md':'current evidence','final-unified-physical-virtual-contract.md':'normative','final-physical-database-architecture.md':'normative','final-access-control-plane-doctrine.md':'normative','api-complete-input-output-contract.md':'normative','three-database-physical-dictionary.md':'normative'}
sources=[]
for p in sorted((ROOT/'docs').glob('*.md'),key=lambda x:x.name.lower()):
    status=status_names.get(p.name,'historical or supporting' if 'handoff' in p.name.lower() or 'plan' in p.name.lower() else 'supporting reference')
    sources.append('<details><summary><b>'+escape(p.name)+'</b><span>'+status+'</span></summary><pre>'+escape(p.read_text(encoding='utf-8',errors='replace'))+'</pre></details>')
ext=ROOT/'documentation'/'complete-package'/'KIDS-admin-panel-complete-source.txt'
if ext.exists(): sources.append('<details><summary><b>DOCX complete extracted source</b><span>preserved source</span></summary><pre>'+escape(ext.read_text(encoding='utf-8',errors='replace'))+'</pre></details>')
css='''*{box-sizing:border-box}:root{--nav:#102c4d;--ink:#182536;--muted:#607086;--line:#d5dee8;--soft:#eef3f8;--blue:#155eef}html{scroll-behavior:smooth}body{margin:0;background:#fbfcfe;color:var(--ink);font:16px/1.72 "Noto Sans Georgian","Sylfaen","Segoe UI",sans-serif}.hero{background:var(--nav);color:white;padding:50px max(26px,calc((100% - 1380px)/2))}.hero h1{font-size:clamp(34px,5vw,58px);line-height:1.08;max-width:900px;margin:0 0 18px}.hero p{max-width:760px;color:#dbe7f3}.layout{max-width:1380px;margin:auto;padding:32px 26px;display:grid;grid-template-columns:270px minmax(0,880px) 190px;gap:34px}nav{position:sticky;top:18px;align-self:start;border-right:1px solid var(--line);padding-right:18px;max-height:94vh;overflow:auto}nav a{display:block;color:#405269;text-decoration:none;padding:7px 0}nav a:hover{color:var(--blue)}main{min-width:0}h2{font-size:31px;line-height:1.2;margin-top:54px}h3{font-size:21px;margin-top:32px}p,li{max-width:78ch}table{width:100%;border-collapse:collapse;margin:20px 0;font-size:14px}th{background:var(--nav);color:white;text-align:left}th,td{border:1px solid var(--line);padding:11px;vertical-align:top}tbody tr:nth-child(even){background:var(--soft)}code{background:#e8eef5;padding:2px 5px}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:white;border:1px solid var(--line);padding:16px;max-height:68vh;overflow:auto;font:12px/1.55 Consolas,monospace}figure{margin:34px 0}figure img{max-width:100%;display:block;border:1px solid var(--line)}figcaption{font-size:13px;color:var(--muted);margin-top:6px}.aside{position:sticky;top:18px;align-self:start;font-size:13px;color:var(--muted)}.aside b{color:var(--ink)}.sources{margin-top:70px}.sources details{border-top:1px solid var(--line)}summary{display:flex;justify-content:space-between;gap:20px;padding:13px 0;cursor:pointer}summary span{color:var(--muted);font-size:13px}.search{width:100%;padding:12px;border:1px solid var(--line);font:inherit}footer{background:var(--nav);color:white;text-align:center;padding:24px}@media(max-width:1050px){.layout{grid-template-columns:230px 1fr}.aside{display:none}}@media(max-width:720px){.layout{display:block}nav{position:static;max-height:none;border-right:0;border-bottom:1px solid var(--line);margin-bottom:28px}}@media(prefers-reduced-motion:reduce){html{scroll-behavior:auto}}'''
for idx,m in enumerate(re.finditer(r'<h2>(.*?)</h2>',guide_html),1): guide_html=guide_html.replace(m.group(0),f'<h2 id="chapter-{idx}">{m.group(1)}</h2>',1)
toc=''.join(f'<a href="#chapter-{i}">{escape(t)}</a>' for i,t in enumerate(re.findall(r'^##\s+(.+)$',guide,re.M),1))+'<a href="#sources">სრული წყაროები</a>'
html='<!doctype html><html lang="ka"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>GEOSOTAT სრული სისტემური დოკუმენტაცია</title><style>'+css+'</style></head><body><header class="hero"><h1>GEOSOTAT სრული სისტემური დოკუმენტაცია</h1><p>ერთი სასწავლო და ნორმატიული ხაზი ბაზებიდან, Control Plane-იდან და Access ingest-იდან API response-მდე.</p></header><div class="layout"><nav><b>სარჩევი</b>'+toc+'</nav><main>'+guide_html+'<h2>არქიტექტურული ნახაზები</h2>'+''.join(images)+'<section class="sources" id="sources"><h2>სრული წყაროების არქივი</h2><p>ყველა Markdown და DOCX extraction სრულად არის ჩასმული. სტატუსი განასხვავებს მიმდინარე, ნორმატიულ და ისტორიულ წყაროს.</p><input id="q" class="search" placeholder="მოძებნე წყაროს ფაილი">'+''.join(sources)+'</section></main><aside class="aside"><b>ფორმულა</b><p>Contract declares<br>Artifact proves<br>Data Plane materializes<br>Snapshot freezes<br>API serves</p><b>მიმდინარე authority</b><p>KIDS R8 status and acceptance</p></aside></div><footer>GEOSOTAT contract driven metadata platform</footer><script>q.oninput=()=>document.querySelectorAll("details").forEach(x=>x.hidden=!x.querySelector("summary").textContent.toLowerCase().includes(q.value.toLowerCase()))</script></body></html>'
(OUT/'START-HERE-geostat-system-complete-documentation.html').write_text(html,encoding='utf-8')

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); pr.append(shd)
def add_table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); tr.append(rep)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'17365D')
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=v; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(cells[i],'F1F5F9' if ri%2 else 'FFFFFF')
    return t

plan=(ROOT/'docs'/'platform-work-and-completion-plan.md').read_text(encoding='utf-8')
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.78); sec.right_margin=Inches(.7)
for n,size in [('Normal',10),('Title',27),('Heading 1',18),('Heading 2',14),('Heading 3',11)]:
    s=doc.styles[n]; s.font.name='Sylfaen'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('GEOSOTAT პლატფორმის სამუშაო და დასრულების გეგმა')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('KIDS R8 მიმდინარე baseline და დარჩენილი platform work').italic=True
doc.add_paragraph('ეს დოკუმენტი არქიტექტურულ დოქტრინას გარდაქმნის შესრულებად სამუშაო პროგრამად. მიმდინარე ფაქტები აღებულია KIDS R8-ის მოქმედი acceptance authority-დან; ძველი handoff ჩანაწერები ისტორიული კონტექსტია.')
doc.add_heading('1 მიმდინარე baseline',1)
add_table(doc,['სფერო','მდგომარეობა'],[('Contract','KIDS_PORTAL_V1 revision 8 APPROVED'),('Ingest','15 dataset შესრულებული'),('Canonical','36 goal, 225 resource, 230 assignment, 178 glossary'),('Statistics','880 observation'),('Publication','snapshot 13 PUBLISHED'),('Serving','297 cache row'),('Archive','2,224 records და pointers'),('Assurance','rollback, replay, quarantine, backup restore და DR')])
doc.add_heading('2 სამუშაოს პრინციპები',1)
for x in ['Contract first','Fail closed','Idempotent execution','Immutable evidence','Schema agnostic runtime','Evidence based completion']: doc.add_paragraph(x,style='List Bullet')
doc.add_heading('3 სამუშაო პაკეტები',1)
rows=[('W01','Documentation authority','ერთი canonical portal და drift check','links/status PASS','P0'),('W02','Management UI','Contract authoring and approval','round trip PASS','P1'),('W03','JWT authorization','OIDC OAuth2 scopes rotation','negative tests PASS','P1'),('W04','Rate quota cost','bounded multi tenant execution','load abuse PASS','P1'),('W05','Metadata negotiation','strict language profile version','matrix PASS','P2'),('W06','SDK generation','Java Kotlin Dart clients','client tests PASS','P2'),('W07','Format providers','SDMX XML Parquet ZIP','round trip PASS','P2'),('W08','High volume','stable key profiles and concurrency','SLO PASS','P2'),('W09','Operations','alerts schedules restore drills','quarterly evidence','P1'),('W10','Legacy retirement','old bindings and paths retired','no consumer','P3')]
add_table(doc,['ID','პაკეტი','შედეგი','Acceptance','Priority'],rows)
doc.add_heading('4 შესრულების თანმიმდევრობა',1)
for x in ['W01 W03 W04 contracts and evidence','W02 UI vertical slice and W09 operations','W05 W06 W07 interoperability','W08 scale and W10 retirement']: doc.add_paragraph(x,style='List Number')
doc.add_heading('5 სავალდებულო evidence',1)
add_table(doc,['Gate','მტკიცებულება'],[('Contract','revision diff და compatibility'),('Database','migration constraints indexes'),('Access','schema hash და round trip'),('Ingest','counts quarantine replay'),('Semantics','metric unit aggregation approvals'),('Privacy','policy და negative test'),('Publication','snapshot cache rollback'),('API','schema relation query errors'),('Operations','health alerts backup DR')])
doc.add_heading('6 Definition of Done',1)
for x in ['ყველა mandatory contract object approved და version consistent არის','Migration repeatable და rollback safe არის','Access approved contract-ს ზუსტად ემთხვევა','Ingest validation quarantine და replay idempotent არის','Quality privacy lineage reconciliation PASS არის','Snapshot immutableა cache შეესაბამება და rollback გამოცდილია','API მხოლოდ declared capabilities-ს ასრულებს','Security quota observability backup და DR evidence შენახულია','Documentation runtime ledger-ს ემთხვევა']: doc.add_paragraph(x,style='List Bullet')
doc.add_heading('7 გადაწყვეტილების კითხვები',1)
for q,a in [('როდის იქმნება ახალი physical table','მხოლოდ ახალი grain-ისას, როცა generic family საკმარისი არ არის და ADR ამას ამტკიცებს.'),('როდის შეიძლება publish','Matching approved revisions და ყველა gate-ის შემდეგ.'),('რას ნიშნავს დასრულება','Deploy, automated test, runtime receipt, reconciliation და rollback evidence.'),('რა არის პირველი შემდეგი ნაბიჯი','W01-ის closure და W03/W04 contract design; შემდეგ UI vertical slice.')]: doc.add_heading(q,2); doc.add_paragraph(a)
doc.add_heading('8 სამუშაო პაკეტების დეტალური კონტრაქტები',1)
details=[
('W01 Documentation authority','მიზანია ერთი canonical ცოდნის წყარო, სადაც current runtime, normative design და historical plans მკაფიოდ არის გამიჯნული.','docs inventory, authority labels, generated HTML, link checker, status drift checker','ყველა 44 Markdown წყარო და DOCX extraction searchable არის; broken link და contradictory current status არ არსებობს.'),
('W02 Management UI','პანელი მართავს contract draft-ს, schema-ს, fields/keys/indexes/relations-ს, semantics-ს, classifier stewardship-ს, gates-ს, publication-სა და audit-ს.','accessible forms, schema preview, compatibility diff, approval workflow, Access generation request, ingest monitor','ერთი ახალი test site კოდის ცვლილების გარეშე გადის draft-იდან published API response-მდე.'),
('W03 JWT and authorization hardening','Production identity უნდა იყოს OIDC/OAuth2-ზე, მოკლე access token-ით, refresh/key rotation-ით და deny-by-default scopes-ით.','issuer/audience validation, JWKS rotation, client/role/scope mapping, audit decision, bootstrap retirement','expired, wrong-audience, wrong-scope და revoked credentials უარყოფილია; ყველა privileged action audit-შია.'),
('W04 Rate quota and cost control','Generic relation/query engine-ს tenant/client budget სჭირდება, რათა complexity და fan-out რესურსს არ გადააჭარბოს.','token bucket, query cost calculator, depth/fan-out/export/concurrency limits, Retry-After and problem details','load/abuse suite ამტკიცებს fairness-ს, predictable latency-ს და fail-closed limit handling-ს.'),
('W05 Metadata negotiation','Locale, response profile და schema revision deterministic negotiation-ით უნდა შეირჩეს.','Accept-Language, Accept-Profile, profile registry, Vary/ETag integration, fallback rules','supported/unsupported locale-profile-version matrix სრულად გადის და cache variants არ ირევა.'),
('W06 SDK generation','OpenAPI/JSON Schema authority-დან უნდა გენერირდეს versioned TypeScript, Java, Kotlin და Dart clients.','generator profiles, package versioning, examples, deprecation markers, CI compile tests','ყველა client compile-დება, example request მუშაობს და breaking revision CI-ში იჭრება.'),
('W07 Format providers','Export core-ს provider SPI სჭირდება SDMX XML, Parquet და ZIP-ისთვის.','format registry, media types, streaming writer, checksum manifest, async artifact lifecycle','round-trip, content negotiation, checksum და large export tests PASS არის.'),
('W08 High volume assurance','ყველა მაღალი მოცულობის family-ს declared stable sort tuple, compatible index და snapshot-bound keyset pagination სჭირდება.','null/order algorithm, forward/back cursor, lastSeen tuple, fingerprint, concurrent insert suite','არც duplicate და არც skipped row; latency და memory SLO განსაზღვრულ ზღვარშია.'),
('W09 Operations hardening','Production ოპერაციები უნდა იყოს observable, schedulable და აღდგენადი.','dashboards, alerts, job schedules, secret rotation, capacity thresholds, backup restore and DR calendar','alert drill, restore drill და runbook replay პერიოდულად ქმნის ხელმოწერილ evidence-ს.'),
('W10 Legacy retirement','Unprefixed bindings, pre-contract endpoints და stale documents კონტროლირებულად უნდა გადავიდეს archive-ში.','usage telemetry, deprecation dates, compatibility adapter expiry, archive manifest','active consumer ძველ path-ზე აღარ არის და removal არ არღვევს published contract-ს.')]
for title,goal,deliver,accept in details:
    doc.add_heading(title,2); doc.add_paragraph('მიზანი. '+goal); doc.add_paragraph('Deliverables. '+deliver); doc.add_paragraph('Acceptance. '+accept)
doc.add_heading('9 Dependencies და critical path',1)
add_table(doc,['წინაპირობა','დამოკიდებული სამუშაო','რატომ'],[('W01','ყველა სამუშაო','source authority და current status უნდა იყოს ერთმნიშვნელოვანი'),('W03 და W04','W02 production activation','პანელი privileged contract actions-ს ასრულებს'),('W05','W06 და public cache','profile/version semantics client-სა და cache-ში უნდა ემთხვეოდეს'),('W07','large export release','format provider checksum და async lifecycle აუცილებელია'),('W08','high-volume onboarding','stable traversal ინდექსისა და snapshot-ის გარეშე ვერ დასტურდება'),('W09','production closure','deploy completion ოპერაციული evidence-ის გარეშე არასრულია')])
doc.add_heading('10 პასუხისმგებლობები',1)
add_table(doc,['როლი','პასუხისმგებლობა','არ შეუძლია მარტო'],[('Business owner','metric, unit, aggregation და publication intent','technical gate-ის გამოტოვება'),('Data steward','classifier, metadata, quality და privacy decisions','source evidence-ის შეცვლა'),('Platform architect','contract boundaries, compatibility და ADR','business meaning-ის გამოგონება'),('Engineer','migration, service, tests და automation','approval-ის თვითნებური მინიჭება'),('Security owner','identity, scopes, secrets და audit policy','data contract-ის ჩუმად შეცვლა'),('Operator','deploy, monitoring, backup, restore და rollback','failed gate-ის გვერდის ავლა')])
doc.add_heading('11 რისკები და კონტროლები',1)
add_table(doc,['რისკი','კონტროლი','Detection'],[('stale documentation','generated inventory და authority labels','CI drift report'),('schema drift','contract hash და preflight','preview failure'),('duplicate replay','idempotency keys და unique indexes','reconciliation'),('semantic fabrication','steward approval gate','unapproved binding query'),('privacy leak','deny-by-default field/profile policy','negative API tests'),('pagination inconsistency','snapshot-bound keyset cursor','concurrent insert suite'),('partial publication','atomic snapshot pointer','member/cache reconciliation'),('unrecoverable artifact','immutable storage, checksum, restore drill','scheduled verification')])
doc.add_heading('12 Release checklist',1)
for x in ['Contract and ingestion revisions match and are approved','All mappings keys relations and indexes validate','Source artifact checksum and object existence verified','Classifier and statistical semantics approved','Quality privacy lineage and reconciliation gates PASS','Snapshot prepared and publication is atomic','Serving cache count reconciles with published members','ETag cursor and response schema tests PASS','Rollback and idempotent replay verified','Monitoring alerts backup restore and DR evidence stored','Documentation status and examples updated']: doc.add_paragraph(x,style='List Bullet')
doc.core_properties.title='GEOSOTAT პლატფორმის სამუშაო და დასრულების გეგმა'; doc.core_properties.author='GEOSOTAT Platform'
doc.save(OUT/'GEOSOTAT-platform-work-and-completion-plan.docx')
print(OUT/'START-HERE-geostat-system-complete-documentation.html'); print(OUT/'GEOSOTAT-platform-work-and-completion-plan.docx')
